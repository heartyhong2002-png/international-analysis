"""
글로벌 IB 리포트 다중 LLM 분석 및 Word(docx) 보고서 생성 파이프라인
1. 유럽계(ING) 리포트는 Mistral-Nemo가 분석
2. 아시아계(Nomura) 리포트는 Qwen2.5가 분석
3. 최종 결과는 EXAONE 3.5가 한국어 임원용 보고서로 종합하여 .docx로 저장
"""

import os
import json
import requests
import pandas as pd
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt
import sys

sys.stdout.reconfigure(encoding='utf-8')

DATA_DIR = Path(__file__).parent.parent / "data" / "signal_gap"
INPUT_CSV = DATA_DIR / "global_ib_insights.csv"
REPORT_DIR = Path(__file__).parent.parent / "reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DOCX = REPORT_DIR / f"Global_IB_Risk_Report_{datetime.now().strftime('%Y%m%d')}.docx"

OLLAMA_URL = "http://localhost:11434/api/generate"

# 기관별 전담 LLM 매핑
IB_LLM_ROUTING = {
    "ING_Think": "mistral-nemo:latest",   # 유럽 시각 (프랑스 Mistral AI)
    "Nomura_Research": "qwen2.5:7b"       # 아시아 시각 (다국어/아시아 특화 Qwen)
}
FINAL_LLM = "exaone3.5:7.8b"              # 최종 보고서 작성 (한국 LG AI EXAONE)

def call_ollama(model_name, prompt, as_json=False):
    payload = {
        "model": model_name,
        "prompt": prompt,
        "stream": False
    }
    if as_json:
        payload["format"] = "json"
        
    try:
        response = requests.post(OLLAMA_URL, json=payload, timeout=300)
        if response.status_code == 200:
            return response.json().get("response", "").strip()
        else:
            return f"[Error] {response.status_code} from Ollama"
    except Exception as e:
        return f"[Error] {e}"

def generate_report():
    print("🧠 다중 LLM 기반 글로벌 IB 리포트 분석 시작...")
    
    if not INPUT_CSV.exists():
        print("⚠️ 수집된 IB 데이터가 없습니다. fetch_ib_research.py를 먼저 실행하세요.")
        return

    df = pd.read_csv(INPUT_CSV)
    
    ib_analyses = {}
    
    # 1단계: 각 기관별 특화 LLM으로 1차 분석
    for inst, group in df.groupby("Institution"):
        if inst not in IB_LLM_ROUTING:
            continue
            
        model = IB_LLM_ROUTING[inst]
        print(f"\n  -> [{inst}] 리포트 분석 중 (전담 모델: {model})...")
        
        texts = "\n".join([f"- {row['Title']}: {row['Summary']}" for _, row in group.iterrows()])
        
        prompt = f"""You are an expert financial analyst. Below are recent research headlines and summaries from {inst}.
Based on these texts, write a concise 1-paragraph summary identifying the core geopolitical or supply chain risks mentioned.

Texts:
{texts}

Summary:"""

        analysis = call_ollama(model, prompt)
        ib_analyses[inst] = analysis
        print(f"     ✓ 1차 분석 완료!")

    # 2단계: EXAONE으로 최종 한국어 종합 보고서 작성
    print(f"\n  -> [최종 종합] 한국어 리포트 작성 중 (전담 모델: {FINAL_LLM})...")
    
    combined_insights = ""
    for inst, text in ib_analyses.items():
        combined_insights += f"\n[Perspective from {inst}]\n{text}\n"

    final_prompt = f"""당신은 한국 최고 수출 대기업의 전략기획실 수석 연구원입니다.
아래에는 유럽계 자본(ING)과 아시아계 자본(Nomura)이 각각 분석한 거시경제 및 지정학 리스크 요약본이 있습니다.

{combined_insights}

이 두 시각을 종합하여, '한국 수출 기업(반도체, 자동차 등)이 대비해야 할 3가지 핵심 공급망 리스크'를 임원진에게 보고하는 형식으로 한국어로 작성해주세요.
반드시 다음 구조를 지켜주세요:
1. 글로벌 자본 시각 요약 (서론)
2. 3대 핵심 리스크 (본론)
3. 전략적 대응 방안 (결론)
"""

    final_report_text = call_ollama(FINAL_LLM, final_prompt)
    print(f"     ✓ EXAONE 3.5 최종 보고서 작성 완료!")

    # 3단계: Word(docx) 파일로 저장
    print("\n📝 MS Word 파일(.docx) 생성 중...")
    doc = Document()
    doc.add_heading('글로벌 IB 지정학 리스크 통합 보고서', 0)
    
    doc.add_paragraph(f"작성일: {datetime.now().strftime('%Y-%m-%d')}")
    doc.add_paragraph(f"분석 소스: ING Think (유럽), Nomura (일본)")
    doc.add_paragraph(f"분석 AI: Mistral-Nemo, Qwen2.5, EXAONE 3.5 (Multi-Agent System)")
    
    doc.add_heading('요약본', level=1)
    doc.add_paragraph(final_report_text)
    
    doc.save(OUTPUT_DOCX)
    print(f"✅ 임원용 보고서 저장 완료! -> {OUTPUT_DOCX}")

if __name__ == "__main__":
    generate_report()

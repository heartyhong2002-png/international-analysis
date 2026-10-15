from pathlib import Path
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
MD = ROOT / 'docs' / 'advisor_meeting' / 'PROJECT_ADVISOR_MEETING_BRIEF.md'
OUT = ROOT / 'docs' / 'advisor_meeting' / 'PROJECT_ADVISOR_MEETING_BRIEF_A4.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def borders(table, color='D9D9D9', size='4'):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_cell_text(cell, text, size=8.2, bold=False, color='000000'):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text.strip())
    r.font.name = 'Malgun Gothic'
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders(table)
    for j, value in enumerate(rows[0]):
        set_cell_text(table.rows[0].cells[j], value, size=8.0, bold=True, color='FFFFFF')
        shade(table.rows[0].cells[j], '1F4E79')
    for row in rows[1:]:
        cells = table.add_row().cells
        for j, value in enumerate(row):
            set_cell_text(cells[j], value, size=7.8)
            if len(table.rows) % 2 == 0:
                shade(cells[j], 'F3F6F9')
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith('|'):
        parts = [p.strip() for p in lines[i].strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-{3,}:?', p) for p in parts):
            i += 1
            continue
        rows.append(parts)
        i += 1
    return rows, i

def add_para(doc, text, style=None, bold_prefix=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    # lightweight markdown emphasis
    chunks = re.split(r'(\*\*.*?\*\*)', text)
    for chunk in chunks:
        if chunk.startswith('**') and chunk.endswith('**'):
            r = p.add_run(chunk[2:-2]); r.bold = True
        else:
            p.add_run(chunk)
    for r in p.runs:
        r.font.name = 'Malgun Gothic'; r.font.size = Pt(9.2)
    return p

def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(16); sec.bottom_margin = Mm(16)
    sec.left_margin = Mm(18); sec.right_margin = Mm(18)
    styles = doc.styles
    normal = styles['Normal']; normal.font.name = 'Malgun Gothic'; normal.font.size = Pt(9.2)
    for name, size, color in [('Title', 20, '000000'), ('Heading 1', 14, '1F4E79'), ('Heading 2', 11, '365F91')]:
        st = styles[name]; st.font.name = 'Malgun Gothic'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    lines = MD.read_text(encoding='utf-8').splitlines()
    i = 0; code = False; code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            if not code:
                code = True; code_lines = []
            else:
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Mm(4); p.paragraph_format.space_after = Pt(5)
                r = p.add_run('\n'.join(code_lines)); r.font.name = 'Consolas'; r.font.size = Pt(7.6)
                code = False
            i += 1; continue
        if code:
            code_lines.append(line); i += 1; continue
        if not line.strip():
            i += 1; continue
        if line.startswith('|'):
            rows, i = parse_table(lines, i); add_table(doc, rows); continue
        m = re.match(r'^(#{1,3})\s+(.*)$', line)
        if m:
            level = len(m.group(1)); text = m.group(2).strip()
            if text == '면담 후 기록란':
                doc.add_page_break()
            p = doc.add_paragraph(text, style='Title' if level == 1 and not doc.paragraphs else f'Heading {level}')
            p.paragraph_format.space_before = Pt(9 if level == 1 else 6); p.paragraph_format.space_after = Pt(5)
            i += 1; continue
        if line.startswith('- [ ] '):
            p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run('☐ ' + line[6:]); r.font.name='Malgun Gothic'; r.font.size=Pt(8.7)
            i += 1; continue
        if line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line[2:]); r.font.name='Malgun Gothic'; r.font.size=Pt(8.8)
            i += 1; continue
        if re.match(r'^\d+\.\s+', line):
            # Keep the explicit number from the Markdown. Word's List Number
            # style continues numbering across sections, which is undesirable
            # for the standalone professor-question list.
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Mm(2); p.paragraph_format.first_line_indent = Mm(-2); p.paragraph_format.space_after = Pt(1)
            r = p.add_run(line); r.font.name='Malgun Gothic'; r.font.size=Pt(8.8)
            i += 1; continue
        add_para(doc, line)
        i += 1
    # footer
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run('국제정세 분석 조기경보 프로그램 · 지도교수 면담 자료')
        run.font.name = 'Malgun Gothic'; run.font.size = Pt(8); run.font.color.rgb = RGBColor(110,110,110)
    doc.save(OUT)
    print(OUT)

if __name__ == '__main__':
    build()
